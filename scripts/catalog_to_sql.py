"""Проверка Excel-каталога и подготовка локального SQL для SQLite/D1."""

from __future__ import annotations

import argparse
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


CATALOG_FIELDS = [
    ("Инвентарный номер", "db_number"),
    ("ID библиогр. записи", "bibliographic_id"),
    ("Штрихкод (852$p)", "inventory_number"),
    ("Состояние фонда", "record_state"),
    ("Тип записи MARC", "record_type"),
    ("Библиогр. уровень", "bibliographic_level"),
    ("Автор (100$a)", "author"),
    ("Заглавие (245$a)", "title"),
    ("Полные сведения 245", "title_full"),
    ("Сведения об издании (250$a)", "edition"),
    ("Место издания (260$a)", "publication_place"),
    ("Издательство (260$b)", "publisher"),
    ("Год издания (260$c)", "publication_year"),
    ("Физическое описание (300$a)", "physical_description"),
    ("Серия (440$a)", "series"),
    ("Темы (650$a)", "subjects"),
    ("Ключевые слова (653$a)", "keywords"),
    ("Классификация (090$a)", "classification"),
    ("Шифр хранения (090$x)", "shelfmark"),
    ("Примечания (952$a)", "notes"),
    ("Местонахождение", "location"),
    ("Статус учёта", "accounting_status"),
    ("Тип фонда", "fund_type"),
    ("Накладная", "invoice"),
    ("Режим инвентаря", "inventory_mode"),
    ("Дата регистрации", "registration_date"),
    ("Дата списания", "writeoff_date"),
    ("Акт списания", "writeoff_act"),
    ("Причина списания", "writeoff_reason"),
    ("Статус выдачи", "loan_status"),
    ("Код читателя", "reader_id"),
    ("Последняя выдача", "last_loan_date"),
    ("Последний возврат", "last_return_date"),
    ("Количество выдач", "loan_count"),
    ("Все поля MARC (JSON)", "marc_fields_json"),
    ("Исходная MARC-запись", "raw_marc"),
]

LOAN_FIELDS = [
    ("Инвентарный номер", "db_number"),
    ("Код читателя", "reader_id"),
    ("Дата выдачи", "loan_date"),
    ("Дата возврата", "return_date"),
    ("Количество", "quantity"),
]


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\x00", "").strip()


def sql(value: Any, integer: bool = False) -> str:
    raw = text(value)
    if integer:
        try:
            return str(max(0, int(float(raw or "0"))))
        except ValueError:
            return "0"
    return "'" + raw.replace("'", "''") + "'"


def sheet_rows(sheet, mapping):
    headers = {text(cell.value): index for index, cell in enumerate(sheet[1])}
    missing = [source for source, _ in mapping if source not in headers]
    if missing:
        raise ValueError(f"На листе {sheet.title!r} отсутствуют поля: {', '.join(missing)}")
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        record = {target: row[headers[source]] for source, target in mapping}
        if any(text(value) for value in record.values()):
            yield row_number, record


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--check-only", action="store_true")
    parser.add_argument("--replace", action="store_true")
    args = parser.parse_args()

    book = load_workbook(args.workbook, read_only=True, data_only=True)
    catalog = list(sheet_rows(book["Каталог"], CATALOG_FIELDS))
    loans = list(sheet_rows(book["Выдачи"], LOAN_FIELDS))
    db_numbers = [text(row["db_number"]) for _, row in catalog]
    nonempty = [value for value in db_numbers if value]
    duplicates = len(nonempty) - len(set(nonempty))
    print(f"Каталог: {len(catalog)}; выдачи: {len(loans)}; пустые № БД: {db_numbers.count('')}; повторы № БД: {duplicates}")
    if args.check_only:
        return
    if not args.output:
        parser.error("для подготовки импорта укажите --output")

    now = datetime.now().isoformat(timespec="seconds")
    lines = ["PRAGMA foreign_keys = ON;", "BEGIN TRANSACTION;"]
    if args.replace:
        lines.extend(["DELETE FROM loans;", "DELETE FROM catalog_records;"])
    catalog_columns = [target for _, target in CATALOG_FIELDS] + ["created_at", "updated_at"]
    for _, record in catalog:
        values = [sql(record[column], column == "loan_count") for column in catalog_columns[:-2]] + [sql(now), sql(now)]
        lines.append(f"INSERT INTO catalog_records ({', '.join(catalog_columns)}) VALUES ({', '.join(values)});")
    loan_columns = [target for _, target in LOAN_FIELDS]
    for _, record in loans:
        values = [sql(record[column], column == "quantity") for column in loan_columns]
        lines.append(f"INSERT INTO loans ({', '.join(loan_columns)}) VALUES ({', '.join(values)});")
    lines.append("COMMIT;")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"SQL подготовлен: {args.output}")


if __name__ == "__main__":
    main()
