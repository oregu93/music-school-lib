from pathlib import Path
from datetime import datetime, date, timezone
import hashlib
import json

from openpyxl import load_workbook


SOURCE = Path("outputs/library_catalog_complete.xlsx")
OUT = Path("import_bundle_final")

SQL_PATH = OUT / "full-import.sql"
MANIFEST_PATH = OUT / "manifest.json"
REPORT_PATH = OUT / "validation.txt"

SYNTHETIC_DB_NUMBER = "LEGACY-MISSING-DB-10924"
REJECT_EXCEL_ROW = 10925


def txt(value):
    if value is None:
        return ""

    if isinstance(value, (datetime, date)):
        return value.isoformat()

    return str(value).strip()


def integer(value):
    if value in (None, ""):
        return 0

    try:
        return int(float(value))
    except (TypeError, ValueError):
        return 0


def sql_value(value):
    if value is None:
        return "NULL"

    if isinstance(value, int):
        return str(value)

    value = str(value)
    return "'" + value.replace("'", "''") + "'"


def sha256(path):
    h = hashlib.sha256()

    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)

    return h.hexdigest()


if not SOURCE.exists():
    raise SystemExit(f"Missing source: {SOURCE}")


OUT.mkdir(exist_ok=True)


wb = load_workbook(
    SOURCE,
    read_only=True,
    data_only=True,
)

catalog = wb["Каталог"]
loans = wb["Выдачи"]

catalog_headers = [txt(c.value) for c in catalog[1]]
loan_headers = [txt(c.value) for c in loans[1]]

ci = {
    name: i
    for i, name in enumerate(catalog_headers)
}

li = {
    name: i
    for i, name in enumerate(loan_headers)
}


catalog_records = []
rejected = []


for excel_row, row in enumerate(
    catalog.iter_rows(
        min_row=2,
        values_only=True,
    ),
    start=2,
):
    if not any(v not in (None, "") for v in row):
        continue

    db_number = txt(
        row[ci["Инвентарный номер"]]
    )

    marc_json = txt(
        row[ci["Все поля MARC (JSON)"]]
    )

    if excel_row == REJECT_EXCEL_ROW:
        rejected.append({
            "excel_row": excel_row,
            "reason": "rejected_empty_source_row",
            "marc_fields_json": marc_json,
        })
        continue

    if not db_number:
        if excel_row != 10924:
            raise RuntimeError(
                f"Unexpected missing db_number at row {excel_row}"
            )

        db_number = SYNTHETIC_DB_NUMBER

    record_state = txt(
        row[ci["Состояние фонда"]]
    )

    source_loan_status = txt(
        row[ci["Статус выдачи"]]
    )

    if record_state == "Списан":
        normalized_loan_status = "Списан"

    elif source_loan_status in (
        "Возвращена",
        "В наличии",
    ):
        normalized_loan_status = "В наличии"

    else:
        raise RuntimeError(
            f"Unexpected loan status at row {excel_row}: "
            f"{source_loan_status!r}"
        )

    catalog_records.append({
        "excel_row": excel_row,
        "db_number": db_number,
        "bibliographic_id":
            txt(row[ci["ID библиогр. записи"]]),
        "inventory_number":
            txt(row[ci["Штрихкод (852$p)"]]),
        "record_state": record_state,
        "record_type":
            txt(row[ci["Тип записи MARC"]]),
        "bibliographic_level":
            txt(row[ci["Библиогр. уровень"]]),
        "author":
            txt(row[ci["Автор (100$a)"]]),
        "title":
            txt(row[ci["Заглавие (245$a)"]]),
        "title_full":
            txt(row[ci["Полные сведения 245"]]),
        "edition":
            txt(row[ci["Сведения об издании (250$a)"]]),
        "publication_place":
            txt(row[ci["Место издания (260$a)"]]),
        "publisher":
            txt(row[ci["Издательство (260$b)"]]),
        "publication_year":
            txt(row[ci["Год издания (260$c)"]]),
        "physical_description":
            txt(row[ci["Физическое описание (300$a)"]]),
        "series":
            txt(row[ci["Серия (440$a)"]]),
        "subjects":
            txt(row[ci["Темы (650$a)"]]),
        "keywords":
            txt(row[ci["Ключевые слова (653$a)"]]),
        "classification":
            txt(row[ci["Классификация (090$a)"]]),
        "shelfmark":
            txt(row[ci["Шифр хранения (090$x)"]]),
        "notes":
            txt(row[ci["Примечания (952$a)"]]),
        "location":
            txt(row[ci["Местонахождение"]]),
        "accounting_status":
            txt(row[ci["Статус учёта"]]),
        "fund_type":
            txt(row[ci["Тип фонда"]]),
        "invoice":
            txt(row[ci["Накладная"]]),
        "inventory_mode":
            txt(row[ci["Режим инвентаря"]]),
        "registration_date":
            txt(row[ci["Дата регистрации"]]),
        "writeoff_date":
            txt(row[ci["Дата списания"]]),
        "writeoff_act":
            txt(row[ci["Акт списания"]]),
        "writeoff_reason":
            txt(row[ci["Причина списания"]]),
        "loan_status": normalized_loan_status,
        "reader_id":
            txt(row[ci["Код читателя"]]),
        "last_loan_date":
            txt(row[ci["Последняя выдача"]]),
        "last_return_date":
            txt(row[ci["Последний возврат"]]),
        "loan_count":
            integer(row[ci["Количество выдач"]]),
        "marc_fields_json": marc_json,
        "raw_marc":
            txt(row[ci["Исходная MARC-запись"]]),
    })


if len(catalog_records) != 10923:
    raise RuntimeError(
        f"Expected 10923 catalog records, got "
        f"{len(catalog_records)}"
    )


db_numbers = [
    r["db_number"]
    for r in catalog_records
]

if len(db_numbers) != len(set(db_numbers)):
    raise RuntimeError(
        "db_number uniqueness failed"
    )


catalog_db_set = set(db_numbers)


loan_records = []


for excel_row, row in enumerate(
    loans.iter_rows(
        min_row=2,
        values_only=True,
    ),
    start=2,
):
    if not any(v not in (None, "") for v in row):
        continue

    db_number = txt(
        row[li["Инвентарный номер"]]
    )

    reader_id = txt(
        row[li["Код читателя"]]
    )

    loan_date = txt(
        row[li["Дата выдачи"]]
    )

    return_date = txt(
        row[li["Дата возврата"]]
    )

    quantity = integer(
        row[li["Количество"]]
    )

    if db_number not in catalog_db_set:
        raise RuntimeError(
            f"Loan row {excel_row} has no catalog parent: "
            f"{db_number!r}"
        )

    if not reader_id:
        raise RuntimeError(
            f"Loan row {excel_row} missing reader_id"
        )

    if not loan_date:
        raise RuntimeError(
            f"Loan row {excel_row} missing loan_date"
        )

    if not return_date:
        raise RuntimeError(
            f"Loan row {excel_row} is unexpectedly open"
        )

    loan_records.append({
        "excel_row": excel_row,
        "db_number": db_number,
        "reader_id": reader_id,
        "loan_date": loan_date,
        "return_date": return_date,
        "quantity": quantity or 1,
    })


if len(loan_records) != 880:
    raise RuntimeError(
        f"Expected 880 loan rows, got "
        f"{len(loan_records)}"
    )


now = datetime.now(
    timezone.utc
).isoformat()


catalog_columns = [
    "db_number",
    "bibliographic_id",
    "inventory_number",
    "record_state",
    "record_type",
    "bibliographic_level",
    "author",
    "title",
    "title_full",
    "edition",
    "publication_place",
    "publisher",
    "publication_year",
    "physical_description",
    "series",
    "subjects",
    "keywords",
    "classification",
    "shelfmark",
    "notes",
    "location",
    "accounting_status",
    "fund_type",
    "invoice",
    "inventory_mode",
    "registration_date",
    "writeoff_date",
    "writeoff_act",
    "writeoff_reason",
    "loan_status",
    "reader_id",
    "last_loan_date",
    "last_return_date",
    "loan_count",
    "marc_fields_json",
    "raw_marc",
    "verified",
    "verified_by",
    "verified_at",
    "deleted_at",
    "deleted_by",
    "created_at",
    "updated_at",
]


with SQL_PATH.open(
    "w",
    encoding="utf-8",
    newline="\n",
) as f:

    f.write("BEGIN TRANSACTION;\n\n")

    for record in catalog_records:

        values = {
            **record,
            "verified": 0,
            "verified_by": None,
            "verified_at": None,
            "deleted_at": None,
            "deleted_by": None,
            "created_at": now,
            "updated_at": now,
        }

        f.write(
            "INSERT INTO catalog_records ("
            + ", ".join(catalog_columns)
            + ") VALUES ("
            + ", ".join(
                sql_value(values[c])
                for c in catalog_columns
            )
            + ");\n"
        )

    f.write("\n")

    for loan in loan_records:

        f.write(
            """
INSERT INTO loans (
    record_id,
    db_number,
    reader_id,
    loan_date,
    return_date,
    quantity,
    reader_note,
    return_note,
    issued_by,
    returned_by,
    created_at
)
SELECT
    id,
    {db_number},
    {reader_id},
    {loan_date},
    {return_date},
    {quantity},
    '',
    '',
    '',
    '',
    {created_at}
FROM catalog_records
WHERE db_number = {db_number};
""".format(
                db_number=sql_value(
                    loan["db_number"]
                ),
                reader_id=sql_value(
                    loan["reader_id"]
                ),
                loan_date=sql_value(
                    loan["loan_date"]
                ),
                return_date=sql_value(
                    loan["return_date"]
                ),
                quantity=loan["quantity"],
                created_at=sql_value(now),
            )
        )

    f.write("\nCOMMIT;\n")


source_sha = sha256(SOURCE)
sql_sha = sha256(SQL_PATH)


manifest = {
    "import_id":
        "IMPORT-02-LEGACY-CATALOG-FINAL",

    "created_at_utc":
        now,

    "source":
        str(SOURCE),

    "source_sha256":
        source_sha,

    "catalog_source_rows":
        10924,

    "catalog_rows_to_insert":
        len(catalog_records),

    "loan_rows_to_insert":
        len(loan_records),

    "synthetic_db_numbers": [
        {
            "excel_row": 10924,
            "assigned_db_number":
                SYNTHETIC_DB_NUMBER,
            "reason":
                "source_db_number_missing",
        }
    ],

    "rejected_rows":
        rejected,

    "normalization": {
        "returned_to_available": 880,
        "written_off_preserved": 1358,
        "verified_default": False,
        "inventory_number_unique_required":
            False,
        "historical_loans_closed": 880,
        "legacy_issued_by": "",
        "legacy_returned_by": "",
    },

    "sql_sha256":
        sql_sha,

    "remote_database_written":
        False,
}


MANIFEST_PATH.write_text(
    json.dumps(
        manifest,
        ensure_ascii=False,
        indent=2,
    ) + "\n",
    encoding="utf-8",
)


report = [
    "IMPORT-02 FINAL VALIDATION",
    "=" * 72,

    f"source_sha256: {source_sha}",
    f"catalog_source_rows: 10924",
    f"catalog_rows_to_insert: "
    f"{len(catalog_records)}",
    f"loan_rows_to_insert: "
    f"{len(loan_records)}",
    "synthetic_db_numbers: 1",
    f"rejected_rows: {len(rejected)}",
    f"unique_db_numbers: "
    f"{len(set(db_numbers))}",
    f"sql_sha256: {sql_sha}",
    "remote_database_written: False",
]


REPORT_PATH.write_text(
    "\n".join(report) + "\n",
    encoding="utf-8",
)


print(
    REPORT_PATH.read_text(
        encoding="utf-8"
    )
)

print("bundle:")
print(" ", SQL_PATH)
print(" ", MANIFEST_PATH)
print(" ", REPORT_PATH)

print()
print("NO D1 DATABASE WAS MODIFIED.")