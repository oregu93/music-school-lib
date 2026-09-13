from pathlib import Path
from collections import Counter, defaultdict
from datetime import datetime, date
import json

from openpyxl import load_workbook


PATH = Path("outputs/library_catalog_complete.xlsx")

wb = load_workbook(
    PATH,
    read_only=True,
    data_only=True,
)

catalog = wb["Каталог"]
loans = wb["Выдачи"]


def txt(value):
    if value is None:
        return ""

    if isinstance(value, (datetime, date)):
        return value.isoformat()

    return str(value).strip()


catalog_headers = [
    txt(cell.value)
    for cell in catalog[1]
]

ci = {
    name: i
    for i, name in enumerate(catalog_headers)
}

loan_headers = [
    txt(cell.value)
    for cell in loans[1]
]

li = {
    name: i
    for i, name in enumerate(loan_headers)
}


records = []

for excel_row, row in enumerate(
    catalog.iter_rows(
        min_row=2,
        values_only=True,
    ),
    start=2,
):
    if not any(v not in (None, "") for v in row):
        continue

    records.append({
        "excel_row": excel_row,
        "db_number":
            txt(row[ci["Инвентарный номер"]]),
        "bibliographic_id":
            txt(row[ci["ID библиогр. записи"]]),
        "inventory_number":
            txt(row[ci["Штрихкод (852$p)"]]),
        "state":
            txt(row[ci["Состояние фонда"]]),
        "loan_status":
            txt(row[ci["Статус выдачи"]]),
        "reader_id":
            txt(row[ci["Код читателя"]]),
        "last_loan":
            txt(row[ci["Последняя выдача"]]),
        "last_return":
            txt(row[ci["Последний возврат"]]),
        "loan_count":
            txt(row[ci["Количество выдач"]]),
        "author":
            txt(row[ci["Автор (100$a)"]]),
        "title":
            txt(row[ci["Заглавие (245$a)"]]),
        "title_full":
            txt(row[ci["Полные сведения 245"]]),
        "marc_json":
            txt(row[ci["Все поля MARC (JSON)"]]),
        "raw_marc":
            txt(row[ci["Исходная MARC-запись"]]),
    })


loan_rows = []

for excel_row, row in enumerate(
    loans.iter_rows(
        min_row=2,
        values_only=True,
    ),
    start=2,
):
    if not any(v not in (None, "") for v in row):
        continue

    loan_rows.append({
        "excel_row": excel_row,
        "db_number":
            txt(row[li["Инвентарный номер"]]),
        "reader_id":
            txt(row[li["Код читателя"]]),
        "loan_date":
            txt(row[li["Дата выдачи"]]),
        "return_date":
            txt(row[li["Дата возврата"]]),
        "quantity":
            txt(row[li["Количество"]]),
    })


print("IMPORT-01C DIAGNOSTICS")
print("=" * 78)


print()
print("1. RECORDS WITHOUT db_number")

missing = [
    r for r in records
    if not r["db_number"]
]

for r in missing:
    print("-" * 78)

    for key, value in r.items():
        if key in ("marc_json", "raw_marc"):
            print(
                f"{key}: "
                f"{'[present]' if value else '[empty]'}"
            )
        else:
            print(f"{key}: {value!r}")

    if r["marc_json"]:
        try:
            parsed = json.loads(r["marc_json"])

            print("MARC JSON preview:")

            if isinstance(parsed, dict):
                for key, value in list(
                    parsed.items()
                )[:20]:
                    print(
                        f"  {key!r}: {value!r}"
                    )
            else:
                print(
                    repr(parsed)[:1500]
                )
        except Exception as exc:
            print(
                "MARC JSON parse error:",
                exc,
            )


print()
print("2. DUPLICATE inventory_number PROFILE")

groups = defaultdict(list)

for r in records:
    if r["inventory_number"]:
        groups[r["inventory_number"]].append(r)

duplicates = {
    key: values
    for key, values in groups.items()
    if len(values) > 1
}

print(
    "distinct duplicated inventory numbers:",
    len(duplicates),
)

print(
    "catalog rows involved:",
    sum(len(v) for v in duplicates.values()),
)

sizes = Counter(
    len(values)
    for values in duplicates.values()
)

print("group-size distribution:")

for size, count in sorted(sizes.items()):
    print(
        f"  {size} records: {count} groups"
    )

print()
print("largest duplicate groups:")

largest = sorted(
    duplicates.items(),
    key=lambda item: (
        -len(item[1]),
        item[0],
    ),
)[:20]

for number, items in largest:
    print()
    print(
        f"{number!r}: {len(items)} records"
    )

    for r in items[:12]:
        print(
            " ",
            r["db_number"],
            "|",
            r["author"][:35],
            "|",
            r["title"][:60],
        )


print()
print("3. LOAN HISTORY COMPLETENESS")

open_loans = [
    r for r in loan_rows
    if not r["return_date"]
]

closed_loans = [
    r for r in loan_rows
    if r["return_date"]
]

missing_loan_date = [
    r for r in loan_rows
    if not r["loan_date"]
]

missing_reader = [
    r for r in loan_rows
    if not r["reader_id"]
]

print(
    "loan rows total:",
    len(loan_rows),
)

print(
    "closed loans:",
    len(closed_loans),
)

print(
    "open loans:",
    len(open_loans),
)

print(
    "missing loan date:",
    len(missing_loan_date),
)

print(
    "missing reader id:",
    len(missing_reader),
)

if open_loans:
    print()
    print("OPEN LOANS:")

    for row in open_loans[:30]:
        print(row)


print()
print("4. CURRENT CATALOG LOAN STATUS")

statuses = Counter(
    r["loan_status"]
    for r in records
)

for status, count in statuses.items():
    print(
        repr(status),
        count,
    )

returned = [
    r for r in records
    if r["loan_status"] == "Возвращена"
]

print()
print(
    "records marked Возвращена:",
    len(returned),
)

print(
    "of them with last_return:",
    sum(
        bool(r["last_return"])
        for r in returned
    ),
)

print(
    "of them with reader_id:",
    sum(
        bool(r["reader_id"])
        for r in returned
    ),
)


print()
print("5. PROPOSED NORMALIZATION")

print(
    "Списан -> Списан:",
    sum(
        r["state"] == "Списан"
        for r in records
    ),
)

print(
    "Возвращена -> В наличии:",
    sum(
        r["loan_status"] == "Возвращена"
        and r["state"] != "Списан"
        for r in records
    ),
)

print(
    "already В наличии:",
    sum(
        r["loan_status"] == "В наличии"
        for r in records
    ),
)

print()
print("NO DATABASE WRITES WERE PERFORMED.")