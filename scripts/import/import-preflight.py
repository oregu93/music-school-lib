from pathlib import Path
from collections import Counter
import json
import sys

from openpyxl import load_workbook

PATH = Path("outputs/library_catalog_complete.xlsx")

if not PATH.exists():
    print(f"ERROR: file not found: {PATH}")
    sys.exit(1)

wb = load_workbook(PATH, read_only=True, data_only=False)

print("IMPORT-01 PREFLIGHT")
print("=" * 72)
print(f"file: {PATH}")
print(f"sheets: {wb.sheetnames}")
print()

for ws in wb.worksheets:
    print(f"SHEET: {ws.title}")
    print(f"max_row: {ws.max_row}")
    print(f"max_column: {ws.max_column}")

    header = [
        ws.cell(row=1, column=i).value
        for i in range(1, ws.max_column + 1)
    ]

    print("headers:")
    for i, value in enumerate(header, start=1):
        print(f"  {i:02d}: {value!r}")

    print()

    if ws.max_row > 1:
        nonempty_rows = 0

        for row in ws.iter_rows(
            min_row=2,
            values_only=True,
        ):
            if any(
                value not in (None, "")
                for value in row
            ):
                nonempty_rows += 1

        print(f"nonempty_data_rows: {nonempty_rows}")

    print("-" * 72)

print()
print("DONE: workbook was opened read-only; no files were modified.")