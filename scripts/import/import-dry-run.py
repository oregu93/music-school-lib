from pathlib import Path
from collections import Counter, defaultdict
from datetime import datetime, date
import json
import sys

from openpyxl import load_workbook


PATH = Path("outputs/library_catalog_complete.xlsx")

CATALOG_SHEET = "Каталог"
LOANS_SHEET = "Выдачи"


EXPECTED_CATALOG_HEADERS = [
    "Инвентарный номер",
    "ID библиогр. записи",
    "Состояние фонда",
    "Тип записи MARC",
    "Библиогр. уровень",
    "Автор (100$a)",
    "Заглавие (245$a)",
    "Полные сведения 245",
    "Сведения об издании (250$a)",
    "Место издания (260$a)",
    "Издательство (260$b)",
    "Год издания (260$c)",
    "Физическое описание (300$a)",
    "Серия (440$a)",
    "Темы (650$a)",
    "Ключевые слова (653$a)",
    "Классификация (090$a)",
    "Шифр хранения (090$x)",
    "Штрихкод (852$p)",
    "Примечания (952$a)",
    "Местонахождение",
    "Статус учёта",
    "Тип фонда",
    "Накладная",
    "Режим инвентаря",
    "Дата регистрации",
    "Дата списания",
    "Акт списания",
    "Причина списания",
    "Статус выдачи",
    "Код читателя",
    "Последняя выдача",
    "Последний возврат",
    "Количество выдач",
    "Все поля MARC (JSON)",
    "Исходная MARC-запись",
]

EXPECTED_LOAN_HEADERS = [
    "Инвентарный номер",
    "Код читателя",
    "Дата выдачи",
    "Дата возврата",
    "Количество",
]


def text(value):
    if value is None:
        return ""

    if isinstance(value, datetime):
        return value.isoformat()

    if isinstance(value, date):
        return value.isoformat()

    return str(value).strip()


def integer(value):
    if value in (None, ""):
        return 0

    try:
        return int(float(value))
    except (TypeError, ValueError):
        return 0


def duplicate_values(values):
    counter = Counter(v for v in values if v)
    return {
        key: count
        for key, count in counter.items()
        if count > 1
    }


if not PATH.exists():
    print(f"ERROR: file not found: {PATH}")
    sys.exit(1)


wb = load_workbook(
    PATH,
    read_only=True,
    data_only=True,
)


if CATALOG_SHEET not in wb.sheetnames:
    raise SystemExit("ERROR: sheet Каталог not found")

if LOANS_SHEET not in wb.sheetnames:
    raise SystemExit("ERROR: sheet Выдачи not found")


catalog_ws = wb[CATALOG_SHEET]
loans_ws = wb[LOANS_SHEET]


catalog_headers = [
    text(cell.value)
    for cell in catalog_ws[1]
]

loan_headers = [
    text(cell.value)
    for cell in loans_ws[1]
]


print("IMPORT-01B DRY RUN")
print("=" * 78)

print()
print("HEADER CHECK")

print(
    "catalog_headers_exact:",
    catalog_headers == EXPECTED_CATALOG_HEADERS,
)

print(
    "loan_headers_exact:",
    loan_headers == EXPECTED_LOAN_HEADERS,
)

if catalog_headers != EXPECTED_CATALOG_HEADERS:
    print("ERROR: catalog header mismatch")
    sys.exit(1)

if loan_headers != EXPECTED_LOAN_HEADERS:
    print("ERROR: loan header mismatch")
    sys.exit(1)


catalog_index = {
    name: i
    for i, name in enumerate(catalog_headers)
}

loan_index = {
    name: i
    for i, name in enumerate(loan_headers)
}


catalog_rows = []

for excel_row_number, row in enumerate(
    catalog_ws.iter_rows(
        min_row=2,
        values_only=True,
    ),
    start=2,
):
    if not any(v not in (None, "") for v in row):
        continue

    record = {
        "excel_row": excel_row_number,

        "db_number":
            text(row[catalog_index["Инвентарный номер"]]),

        "bibliographic_id":
            text(row[catalog_index["ID библиогр. записи"]]),

        "record_state":
            text(row[catalog_index["Состояние фонда"]]),

        "record_type":
            text(row[catalog_index["Тип записи MARC"]]),

        "bibliographic_level":
            text(row[catalog_index["Библиогр. уровень"]]),

        "author":
            text(row[catalog_index["Автор (100$a)"]]),

        "title":
            text(row[catalog_index["Заглавие (245$a)"]]),

        "title_full":
            text(row[catalog_index["Полные сведения 245"]]),

        "edition":
            text(row[catalog_index[
                "Сведения об издании (250$a)"
            ]]),

        "publication_place":
            text(row[catalog_index[
                "Место издания (260$a)"
            ]]),

        "publisher":
            text(row[catalog_index[
                "Издательство (260$b)"
            ]]),

        "publication_year":
            text(row[catalog_index[
                "Год издания (260$c)"
            ]]),

        "physical_description":
            text(row[catalog_index[
                "Физическое описание (300$a)"
            ]]),

        "series":
            text(row[catalog_index["Серия (440$a)"]]),

        "subjects":
            text(row[catalog_index["Темы (650$a)"]]),

        "keywords":
            text(row[catalog_index[
                "Ключевые слова (653$a)"
            ]]),

        "classification":
            text(row[catalog_index[
                "Классификация (090$a)"
            ]]),

        "shelfmark":
            text(row[catalog_index[
                "Шифр хранения (090$x)"
            ]]),

        "inventory_number":
            text(row[catalog_index["Штрихкод (852$p)"]]),

        "notes":
            text(row[catalog_index[
                "Примечания (952$a)"
            ]]),

        "location":
            text(row[catalog_index["Местонахождение"]]),

        "accounting_status":
            text(row[catalog_index["Статус учёта"]]),

        "fund_type":
            text(row[catalog_index["Тип фонда"]]),

        "invoice":
            text(row[catalog_index["Накладная"]]),

        "inventory_mode":
            text(row[catalog_index["Режим инвентаря"]]),

        "registration_date":
            text(row[catalog_index["Дата регистрации"]]),

        "writeoff_date":
            text(row[catalog_index["Дата списания"]]),

        "writeoff_act":
            text(row[catalog_index["Акт списания"]]),

        "writeoff_reason":
            text(row[catalog_index["Причина списания"]]),

        "loan_status":
            text(row[catalog_index["Статус выдачи"]]),

        "reader_id":
            text(row[catalog_index["Код читателя"]]),

        "last_loan_date":
            text(row[catalog_index["Последняя выдача"]]),

        "last_return_date":
            text(row[catalog_index["Последний возврат"]]),

        "loan_count":
            integer(row[catalog_index["Количество выдач"]]),

        "marc_fields_json":
            text(row[catalog_index[
                "Все поля MARC (JSON)"
            ]]),

        "raw_marc":
            text(row[catalog_index[
                "Исходная MARC-запись"
            ]]),
    }

    catalog_rows.append(record)


loan_rows = []

for excel_row_number, row in enumerate(
    loans_ws.iter_rows(
        min_row=2,
        values_only=True,
    ),
    start=2,
):
    if not any(v not in (None, "") for v in row):
        continue

    loan_rows.append({
        "excel_row": excel_row_number,
        "db_number":
            text(row[loan_index["Инвентарный номер"]]),
        "reader_id":
            text(row[loan_index["Код читателя"]]),
        "loan_date":
            text(row[loan_index["Дата выдачи"]]),
        "return_date":
            text(row[loan_index["Дата возврата"]]),
        "quantity":
            integer(row[loan_index["Количество"]]),
    })


print()
print("ROW COUNTS")
print("catalog_rows:", len(catalog_rows))
print("loan_rows:", len(loan_rows))


db_numbers = [
    r["db_number"]
    for r in catalog_rows
]

inventory_numbers = [
    r["inventory_number"]
    for r in catalog_rows
]


empty_db = [
    r for r in catalog_rows
    if not r["db_number"]
]

empty_inventory = [
    r for r in catalog_rows
    if not r["inventory_number"]
]

dup_db = duplicate_values(db_numbers)
dup_inventory = duplicate_values(inventory_numbers)


print()
print("IDENTITY CHECKS")

print("empty_db_number:", len(empty_db))
print("duplicate_db_number_values:", len(dup_db))

print("empty_inventory_number:", len(empty_inventory))
print(
    "duplicate_inventory_number_values:",
    len(dup_inventory),
)


states = Counter(
    r["record_state"]
    for r in catalog_rows
)

loan_statuses = Counter(
    r["loan_status"]
    for r in catalog_rows
)

print()
print("CATALOG STATES")

for key, count in states.most_common():
    print(repr(key), count)


print()
print("LOAN STATUS VALUES")

for key, count in loan_statuses.most_common():
    print(repr(key), count)


invalid_json = []

for record in catalog_rows:
    value = record["marc_fields_json"]

    if not value:
        continue

    try:
        json.loads(value)
    except Exception as exc:
        invalid_json.append(
            (
                record["excel_row"],
                record["db_number"],
                str(exc),
            )
        )


print()
print("MARC CHECKS")

print(
    "nonempty_marc_json:",
    sum(
        bool(r["marc_fields_json"])
        for r in catalog_rows
    ),
)

print("invalid_marc_json:", len(invalid_json))

print(
    "nonempty_raw_marc:",
    sum(
        bool(r["raw_marc"])
        for r in catalog_rows
    ),
)


catalog_db_set = {
    value
    for value in db_numbers
    if value
}

loan_without_catalog = [
    row
    for row in loan_rows
    if row["db_number"] not in catalog_db_set
]


print()
print("LOAN RELATION CHECK")

print(
    "loan_rows_without_catalog_record:",
    len(loan_without_catalog),
)


loan_db_counts = Counter(
    row["db_number"]
    for row in loan_rows
)

print(
    "distinct_catalog_records_with_loan_history:",
    len([
        key
        for key in loan_db_counts
        if key in catalog_db_set
    ]),
)


print()
print("IMPORTANT FIELD COVERAGE")

for key in [
    "db_number",
    "bibliographic_id",
    "inventory_number",
    "author",
    "title",
    "publication_year",
    "record_state",
    "loan_status",
    "marc_fields_json",
    "raw_marc",
]:
    count = sum(
        bool(record[key])
        for record in catalog_rows
    )

    print(
        f"{key}: {count}/{len(catalog_rows)}"
    )


print()
print("SAMPLE PROBLEM ROWS")

if empty_db:
    print("empty db_number:")
    for row in empty_db[:10]:
        print(
            row["excel_row"],
            row["title"],
            row["inventory_number"],
        )

if dup_db:
    print("duplicate db_number:")
    for value, count in list(dup_db.items())[:20]:
        print(repr(value), count)

if dup_inventory:
    print("duplicate inventory_number:")
    for value, count in list(
        dup_inventory.items()
    )[:20]:
        print(repr(value), count)

if loan_without_catalog:
    print("loan rows without catalog:")
    for row in loan_without_catalog[:20]:
        print(row)

if invalid_json:
    print("invalid MARC JSON:")
    for row in invalid_json[:10]:
        print(row)


print()
print("DRY-RUN IMPORT PROJECTION")

print(
    "catalog_records_to_insert:",
    len(catalog_rows),
)

print(
    "loan_rows_to_insert:",
    len(loan_rows)
)

print(
    "catalog_rows_blocked_by_missing_db_number:",
    len(empty_db),
)

print(
    "catalog_rows_with_duplicate_db_number:",
    sum(dup_db.values()),
)

print(
    "loan_rows_without_parent:",
    len(loan_without_catalog),
)


safe_catalog = (
    len(catalog_rows)
    - len(empty_db)
)

print(
    "candidate_catalog_rows_after_required-key_check:",
    safe_catalog,
)


print()
print("NO DATABASE WRITES WERE PERFORMED.")