from pathlib import Path
import hashlib
import json
import unicodedata

from openpyxl import load_workbook


SOURCE = Path(
    "outputs/library_catalog_complete.xlsx"
)

OUT_DIR = Path("search_01")
OUT_SQL = OUT_DIR / "backfill-search-text.sql"
OUT_REPORT = OUT_DIR / "validation.txt"

SYNTHETIC_DB_NUMBER = (
    "LEGACY-MISSING-DB-10924"
)

REJECT_ROW = 10925


def text(value):
    if value is None:
        return ""
    return str(value).strip()


def normalize(value):
    return unicodedata.normalize(
        "NFKC",
        text(value),
    ).lower()


def sql_string(value):
    return "'" + str(value).replace(
        "'",
        "''",
    ) + "'"


def sha256(path):
    h = hashlib.sha256()

    with path.open("rb") as f:
        for chunk in iter(
            lambda: f.read(1024 * 1024),
            b"",
        ):
            h.update(chunk)

    return h.hexdigest()


OUT_DIR.mkdir(exist_ok=True)

wb = load_workbook(
    SOURCE,
    read_only=True,
    data_only=True,
)

ws = wb["Каталог"]

headers = [
    text(cell.value)
    for cell in ws[1]
]

idx = {
    name: i
    for i, name in enumerate(headers)
}


updates = []


for excel_row, row in enumerate(
    ws.iter_rows(
        min_row=2,
        values_only=True,
    ),
    start=2,
):
    if not any(
        value not in (None, "")
        for value in row
    ):
        continue

    if excel_row == REJECT_ROW:
        continue

    db_number = text(
        row[idx["Инвентарный номер"]]
    )

    if not db_number:
        if excel_row != 10924:
            raise RuntimeError(
                "Unexpected missing db_number "
                f"at row {excel_row}"
            )

        db_number = SYNTHETIC_DB_NUMBER

    fields = {
        "dbNumber":
            normalize(
                db_number
            ),

        "inventoryNumber":
            normalize(
                row[idx["Штрихкод (852$p)"]]
            ),

        "author":
            normalize(
                row[idx["Автор (100$a)"]]
            ),

        "title":
            normalize(
                row[idx["Заглавие (245$a)"]]
            ),

        "titleFull":
            normalize(
                row[idx["Полные сведения 245"]]
            ),

        "publisher":
            normalize(
                row[idx["Издательство (260$b)"]]
            ),

        "year":
            normalize(
                row[idx["Год издания (260$c)"]]
            ),

        "subjects":
            normalize(
                row[idx["Темы (650$a)"]]
            ),

        "keywords":
            normalize(
                row[idx["Ключевые слова (653$a)"]]
            ),

        "shelfmark":
            normalize(
                row[idx["Шифр хранения (090$x)"]]
            ),

        "notes":
            normalize(
                row[idx["Примечания (952$a)"]]
            ),
    }

    search = {
        **fields,
        "all": "\n".join(
            value
            for value in fields.values()
            if value
        ),
    }

    encoded = json.dumps(
        search,
        ensure_ascii=False,
        separators=(",", ":"),
    )

    updates.append(
        "UPDATE catalog_records "
        f"SET search_text={sql_string(encoded)} "
        f"WHERE db_number={sql_string(db_number)};"
    )


if len(updates) != 10923:
    raise RuntimeError(
        f"Expected 10923 updates, got "
        f"{len(updates)}"
    )


OUT_SQL.write_text(
    "\n".join(updates) + "\n",
    encoding="utf-8",
)


report = [
    "SEARCH-01 BACKFILL",
    "=" * 72,
    f"source_sha256: {sha256(SOURCE)}",
    f"rows_to_update: {len(updates)}",
    f"sql_sha256: {sha256(OUT_SQL)}",
    "database_written: False",
]


OUT_REPORT.write_text(
    "\n".join(report) + "\n",
    encoding="utf-8",
)


print(
    OUT_REPORT.read_text(
        encoding="utf-8"
    )
)